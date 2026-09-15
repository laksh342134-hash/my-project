from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from report import generate_report


def export_excel_report(report_df, output_path="risk_report.xlsx"):
    report_df.to_excel(output_path, index=False)

    workbook = load_workbook(output_path)
    worksheet = workbook.active
    colors = {"High": "FFC7CE", "Medium": "FFEB9C", "Low": "C6EFCE"}

    for row in range(2, worksheet.max_row + 1):
        risk_cell = worksheet.cell(row=row, column=4)
        fill_color = colors.get(risk_cell.value, "FFFFFF")
        risk_cell.fill = PatternFill(
            start_color=fill_color,
            end_color=fill_color,
            fill_type="solid",
        )

    workbook.save(output_path)


if __name__ == "__main__":
    report = generate_report("sample_data.csv")
    export_excel_report(report)
    print("Excel report created: risk_report.xlsx")
